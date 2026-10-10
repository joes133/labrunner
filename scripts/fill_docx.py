#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
fill_docx.py - 按占位符填充实验报告模板

核心原则：
    - 只改 run.text，不动 run.font，保留原模板的全部格式
    - 禁止正则替换 docx 内部 XML（那是乱码与文档损坏的源头）
    - 写完立即回读校验

占位符格式：
    {{学号}}  {{姓名}}  {{实验目的}}      文本
    {{img:step-01-编译.png}}              图片（配合 --images 目录）

用法：
    python fill_docx.py --template 模板.docx --list       列出所有占位符
    python fill_docx.py --template t.docx --data d.json --output out.docx
    python fill_docx.py --template t.docx --data '{"学号":"123"}' \
                        --images ./shots --output out.docx

依赖：python-docx
"""

import argparse
import json
import re
import sys
from pathlib import Path

# Windows 终端默认可能是 GBK，强制 UTF-8 输出，避免中文提示变乱码
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

TOKEN_RE = re.compile(r"\{\{([^}]+)\}\}")
IMG_PREFIX = "img:"

try:
    import docx
    from docx.shared import Inches
except ImportError:
    docx = None


def require_docx():
    if docx is None:
        print("缺少 python-docx，请安装：pip install python-docx", file=sys.stderr)
        print("不要用 zipfile + 正则手工改 docx —— 那是乱码的源头", file=sys.stderr)
        sys.exit(1)


def iter_paragraphs(doc):
    """遍历正文 + 页眉 + 页脚 + 表格单元格里的所有段落"""
    for p in doc.paragraphs:
        yield p
    for section in doc.sections:
        for p in section.header.paragraphs:
            yield p
        for p in section.footer.paragraphs:
            yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def collect_tokens(doc):
    found = []
    for p in iter_paragraphs(doc):
        for m in TOKEN_RE.finditer(p.text):
            found.append(m.group(1))
    return sorted(set(found))


def replace_token_in_paragraph(p, token, value):
    """在单个段落内替换 {{token}}，尽量保留 run 格式。成功返回 True。"""
    full = p.text
    if token not in full:
        return False

    idx = full.find(token)
    end = idx + len(token)
    pos = 0
    placed = False

    for run in p.runs:
        rstart = pos
        rend = pos + len(run.text)
        pos = rend

        # 无交集
        if rend <= idx or rstart >= end:
            continue

        # 有交集
        if not placed:
            before = run.text[: max(0, idx - rstart)]
            after = run.text[max(0, end - rstart) :]
            run.text = before + value + after
            placed = True
        else:
            # token 的剩余部分落在这个 run 里，删掉
            cut_start = max(0, idx - rstart)
            cut_end = min(len(run.text), end - rstart)
            run.text = run.text[:cut_start] + run.text[cut_end:]

    return placed


def replace_all_in_paragraph(p, token, value, max_iter=20):
    """替换段落内所有 token，返回替换次数。

    注意 token 必须带花括号，例如 {{学号}}。
    不带花括号会只替换内部文字，留下空壳花括号。
    """
    count = 0
    while token in p.text and count < max_iter:
        if not replace_token_in_paragraph(p, token, value):
            break
        count += 1
    return count


def insert_image_after(p, path, width_in=5.5):
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))


def fill(doc, data, images_dir=None, width_in=5.5):
    """填充所有占位符，返回统计"""
    stats = {"text": 0, "image": 0, "missing": []}

    # 先处理图片占位符（会改变段落内容）
    for p in list(iter_paragraphs(doc)):
        for key in TOKEN_RE.findall(p.text):
            if not key.startswith(IMG_PREFIX):
                continue
            name = key[len(IMG_PREFIX) :].strip()
            if images_dir:
                path = Path(images_dir) / name
            else:
                path = Path(name)
            if not path.exists():
                stats["missing"].append(name)
                continue
            # 占位符必须带花括号一起替换，否则会留下 {{}} 空壳
            replace_all_in_paragraph(p, "{{" + key + "}}", "")
            insert_image_after(p, path, width_in)
            stats["image"] += 1

    # 再处理文本占位符
    for key, value in data.items():
        token = "{{" + key + "}}"
        for p in list(iter_paragraphs(doc)):
            stats["text"] += replace_all_in_paragraph(p, token, str(value))

    return stats


def verify(path, expect_keys=()):
    """回读校验：占位符是否清干净、有没有乱码、关键字段是否写入"""
    d = docx.Document(str(path))
    text = "\n".join(p.text for p in iter_paragraphs(d))
    problems = []

    left = TOKEN_RE.findall(text)
    if left:
        problems.append("仍有未替换的占位符: {}".format(sorted(set(left))))

    if "\ufffd" in text:
        problems.append("检测到 U+FFFD 替换字符，编码可能已损坏")

    for k in expect_keys:
        if k and k not in text:
            problems.append("关键内容未写入: {}".format(k))

    return problems


def load_data(spec):
    p = Path(spec)
    if p.exists():
        return json.loads(p.read_text(encoding="utf-8"))
    return json.loads(spec)


def main():
    ap = argparse.ArgumentParser(description="按占位符填充实验报告模板")
    ap.add_argument("--template", required=True, help="模板 docx 路径")
    ap.add_argument("--output", help="输出路径（--list 时不需要）")
    ap.add_argument("--data", help="填充数据：JSON 文件路径或 JSON 字符串")
    ap.add_argument("--images", help="图片目录，配合 {{img:xxx.png}} 使用")
    ap.add_argument("--width", type=float, default=5.5, help="图片宽度（英寸），默认 5.5")
    ap.add_argument("--list", action="store_true", help="只列出模板中的占位符")
    ap.add_argument("--verify", action="store_true", help="写完后回读校验")
    args = ap.parse_args()

    require_docx()

    tpl = Path(args.template)
    if not tpl.exists():
        print("模板不存在: {}".format(tpl), file=sys.stderr)
        return 1

    doc = docx.Document(str(tpl))

    if args.list:
        tokens = collect_tokens(doc)
        if not tokens:
            print("模板中没有找到占位符")
        else:
            print("模板占位符共 {} 个：".format(len(tokens)))
            for t in tokens:
                print("  {}{}{}".format("{{", t, "}}"))
        return 0

    if not args.data:
        print("需要 --data 提供填充内容", file=sys.stderr)
        return 1
    if not args.output:
        print("需要 --output 指定输出路径", file=sys.stderr)
        return 1

    data = load_data(args.data)
    stats = fill(doc, data, args.images, args.width)

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))

    print("已生成: {}".format(out))
    print("  文本替换 {} 处，图片插入 {} 张".format(stats["text"], stats["image"]))
    if stats["missing"]:
        print("  警告: 以下图片未找到 -> {}".format(stats["missing"]), file=sys.stderr)

    if args.verify:
        problems = verify(out, expect_keys=[str(v) for v in data.values()])
        if problems:
            print("校验未通过:", file=sys.stderr)
            for p in problems:
                print("  - {}".format(p), file=sys.stderr)
            return 2
        print("  校验通过：无残留占位符、无乱码、关键内容已写入")

    return 0


if __name__ == "__main__":
    sys.exit(main())
